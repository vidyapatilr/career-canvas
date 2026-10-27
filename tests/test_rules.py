import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import io

import docx

from core.fit_review import label_for
from core.resume_builder import education_for, to_docx
from core.rules import check_cover_letter, check_resume
from core import tracker_db


def test_label_boundaries():
    assert [label_for(s) for s in (85, 84, 65, 64, 45, 44)] == ["Perfect", "Good", "Good", "Stretch", "Stretch", "Skip"]


def test_ccra_only_for_healthcare():
    assert any("CCRA" in e for e in education_for(True))
    assert not any("CCRA" in e for e in education_for(False))


def _resume(bullets):
    return {"summary": "QA leader focused on data quality.", "competencies": ["ETL testing"],
            "jobs": [{"title": "QA — Hitachi Vantara", "dates": "Nov 2018 – Jul 2026", "bullets": bullets}]}


def test_resume_checks():
    good = [f"Validated pipeline number {i}." for i in range(6)]
    assert check_resume(_resume(good), [], []) == []
    bad = good[:5] + ["Built TestNG suites — fast", "Has 20+ years in low code tools"]
    problems = " ".join(check_resume(_resume(bad), [], []))
    for needle in ("TestNG", "Dash", "years", "low code"):
        assert needle in problems
    assert "max 7" not in problems
    assert "max 7" in " ".join(check_resume(_resume(good + good[:2]), [], []))


def test_hyphenated_words_ok():
    assert check_resume(_resume([f"Automated Docker-based setup {i}." for i in range(6)]), [], []) == []


def test_cover_letter_checks():
    ok = {"re_line": "Re: QA Lead", "paragraphs": ["One.", "Two.", "Three."]}
    assert check_cover_letter(ok) == []
    assert check_cover_letter({"re_line": "Re: QA Lead", "paragraphs": ["FindMyTrial was great."]})


def test_docx_fonts_and_ccra():
    data = _resume([f"Validated pipeline {i}." for i in range(6)])
    data["education"] = education_for(False)
    d = docx.Document(io.BytesIO(to_docx(data)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "CCRA" not in text
    name_run = d.paragraphs[0].runs[0]
    assert name_run.bold and name_run.font.size.pt == 16
    assert d.styles["Normal"].font.name == "Calibri"


def test_tracker_crud(tmp_path):
    db = tmp_path / "t.db"
    rid = tracker_db.add("Acme", "QA Lead", "2026-10-27", db_path=db)
    tracker_db.update(rid, status="Interview", db_path=db)
    assert tracker_db.list_all(db)[0]["status"] == "Interview"
    tracker_db.delete(rid, db_path=db)
    assert tracker_db.list_all(db) == []
